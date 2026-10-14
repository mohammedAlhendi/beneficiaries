"""تصدير الجداول إلى إكسل بتخطيط من اليمين إلى اليسار (تدفقي، يتحمل مئات آلاف الصفوف)."""
from django.http import HttpResponse
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _width(header):
    h = str(header)
    if "الاسم" in h or "الحالة" in h or "السبب" in h:
        return 34
    return max(len(h) + 4, 14)


def xlsx_response(filename, title, headers, rows, subtitle=""):
    """rows: أي iterable من الصفوف (قائمة أو مولّد)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=(title or "كشف")[:31])
    ws.sheet_view.rightToLeft = True
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _width(h)

    top = WriteOnlyCell(ws, value=title)
    top.font = Font(bold=True, size=14)
    ws.append([top])
    header_row = 2
    if subtitle:
        ws.append([subtitle])
        header_row += 1
    ws.append([])
    header_row += 1
    ws.freeze_panes = f"A{header_row + 1}"

    head_cells = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5C4A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        head_cells.append(c)
    ws.append(head_cells)

    for row in rows:
        ws.append(list(row))

    response = HttpResponse(content_type=XLSX_TYPE)
    response["Content-Disposition"] = content_disposition_header(True, filename)
    wb.save(response)
    return response
